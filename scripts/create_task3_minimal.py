"""Gera o notebook mínimo da atividade 3 e seu enunciado em Word.

O texto do enunciado é mantido em tarefas/Tarefa_3_..._Versao_Minima.md.
Requer python-docx. Execute: python scripts/create_task3_minimal.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from textwrap import dedent
from urllib.parse import quote, unquote

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from create_task1_docx import (
    TEAL,
    add_hyperlink,
    add_page_number,
    add_text,
    configure_section,
    style_document,
)


ROOT = Path(__file__).resolve().parents[1]
NOTEBOOK = "03_aprendizado_nao_supervisionado_versão_minima.ipynb"
TASK = ROOT / "tarefas" / "Tarefa_3_Aprendizado_Nao_Supervisionado_Versao_Minima.md"
REPO_URL = "https://github.com/flavioluizseixas/aprendizado-de-maquina-para-saude"
COLAB = REPO_URL.replace("github.com", "colab.research.google.com/github")


def cell(kind: str, source: str) -> dict:
    result = {
        "cell_type": kind,
        "metadata": {},
        "source": dedent(source).strip().splitlines(keepends=True),
    }
    if kind == "code":
        result.update(execution_count=None, outputs=[])
    return result


def build_notebook() -> dict:
    cells = [
        cell("markdown", f"""
            # 03 — Aprendizado não supervisionado: versão mínima

            [![Abrir no Colab](https://colab.research.google.com/assets/colab-badge.svg)]({COLAB}/blob/main/notebooks/{quote(NOTEBOOK)})

            **Objetivo:** formar e interpretar grupos com K-means e visualizá-los com PCA. Execute as seis células de código em ordem. Reserve 40–50 minutos para a atividade completa.

            Use o Colab ou Jupyter com `pandas`, `matplotlib` e `scikit-learn` instalados. Se necessário, execute `%pip install pandas matplotlib scikit-learn` em uma célula adicional. O carregamento requer internet; o notebook funciona sem clonar o repositório. A primeira célula informa as etapas de importação e download; a conexão usa timeout de 30 segundos por operação de rede. Se nenhuma mensagem aparecer, reinicie o kernel e confirme o ambiente Python selecionado.

            ## Fonte e licença

            [CDC Diabetes Health Indicators — UCI 891](https://archive.ics.uci.edu/dataset/891/cdc+diabetes+health+indicators). Consulte a fonte para o dicionário e os termos de uso. Cada linha representa uma pessoa.

            | Atributo | Interpretação |
            |---|---|
            | `BMI` | IMC em kg/m² |
            | `Age` | Faixa etária (código 1–13, não idade em anos) |
            | `GenHlth` | Saúde geral: 1 = excelente; 5 = ruim |
            | `PhysHlth`, `MentHlth` | Dias de saúde física/mental ruim nos últimos 30 dias |
            | `Education`, `Income` | Escolaridade (1–6) e renda (1–8), em categorias ordenadas |

            Tratar os códigos ordinais como números é uma simplificação das distâncias entre categorias. `ID` identifica o registro. **`Diabetes_binary` não entra no agrupamento nem no PCA:** 0 = sem diabetes; 1 = pré-diabetes/diabetes. Será consultado somente ao descrever os grupos já formados.

            Material com finalidade exclusivamente educacional. Os dados incluem autorrelato; os grupos não são diagnósticos e os percentuais descrevem somente esta amostra.

            ## 1. Carregar e selecionar os dados

            Usaremos uma amostra aleatória de 5.000 registros, com semente 42, para uma execução leve. Excluímos eventuais ausências apenas nos atributos usados. A análise é exploratória nesta amostra, sem avaliação preditiva em um conjunto de teste.
        """),
        cell("code", """
            print("1/3 — Importando bibliotecas...", flush=True)
            from io import BytesIO
            from urllib.error import URLError
            from urllib.request import urlopen

            import pandas as pd
            import matplotlib.pyplot as plt
            from IPython.display import display
            from sklearn.preprocessing import StandardScaler
            from sklearn.cluster import KMeans
            from sklearn.metrics import silhouette_score
            from sklearn.decomposition import PCA

            atributos = ["BMI", "Age", "GenHlth", "PhysHlth", "MentHlth", "Education", "Income"]
            print("2/3 — Baixando dados da UCI (~13 MB)...", flush=True)
            url = "https://archive.ics.uci.edu/static/public/891/data.csv"
            try:
                with urlopen(url, timeout=30) as resposta:
                    conteudo = resposta.read()
            except (URLError, TimeoutError, ConnectionError) as erro:
                raise RuntimeError(
                    "Não foi possível baixar os dados da UCI. Verifique a conexão "
                    "com a internet e execute esta célula novamente."
                ) from erro

            print("3/3 — Lendo o CSV e selecionando a amostra...", flush=True)
            dados = pd.read_csv(BytesIO(conteudo), index_col="ID")
            display(dados[atributos].isna().sum().to_frame("Ausências na base"))
            dados = dados.dropna(subset=atributos).sample(n=5_000, random_state=42).copy()
            X = dados[atributos]
            print(f"Amostra: {len(X):,} pessoas | Atributos: {X.shape[1]}")
            display(X.head())
        """),
        cell("markdown", """
            ## 2. Padronizar os atributos

            O `StandardScaler` subtrai a média e divide pelo desvio-padrão de cada atributo. Assim, diferenças de unidade não dominam as distâncias. Isso não elimina a influência de valores extremos nem resolve as limitações dos códigos ordinais.
        """),
        cell("code", """
            padronizador = StandardScaler()
            X_pad = padronizador.fit_transform(X)
        """),
        cell("markdown", """
            ## 3. Comparar números de clusters

            **Inércia** resume as distâncias quadráticas aos centros: procure um cotovelo, onde a redução fica menos acentuada. **Silhouette** varia de −1 a 1; valores maiores indicam grupos mais compactos e separados, e valores próximos de zero sugerem sobreposição. Ele será calculado na mesma subamostra de 2.000 registros para cada k, para reduzir o custo.

            O K-means usa as sete variáveis padronizadas, com dez inicializações por k. Nenhum dos critérios garante que existam grupos clinicamente relevantes.
        """),
        cell("code", """
            resultados = []
            for k in range(2, 7):
                modelo = KMeans(n_clusters=k, n_init=10, random_state=42)
                grupos = modelo.fit_predict(X_pad)
                silhouette = silhouette_score(
                    X_pad, grupos, sample_size=2_000, random_state=42
                )
                resultados.append({"k": k, "Inércia": modelo.inertia_, "Silhouette": silhouette})

            resultados = pd.DataFrame(resultados)
            display(resultados.round(3))
            fig, eixos = plt.subplots(1, 2, figsize=(9, 3))
            for eixo, metrica in zip(eixos, ["Inércia", "Silhouette"]):
                eixo.plot(resultados["k"], resultados[metrica], marker="o")
                eixo.set(xlabel="Número de clusters (k)", ylabel=metrica, xticks=range(2, 7))
            plt.tight_layout()
            plt.show()
        """),
        cell("markdown", """
            ## 4. Escolher k e formar os grupos

            Começamos pelo maior silhouette entre os valores testados. Compare com o cotovelo e justifique a decisão; para testar outra escolha, substitua a primeira linha por `k_escolhido = 3`, por exemplo. Não use `Diabetes_binary` para escolher k. Os rótulos 0, 1, … não indicam ordem de gravidade.
        """),
        cell("code", """
            k_escolhido = int(resultados.loc[resultados["Silhouette"].idxmax(), "k"])
            kmeans = KMeans(n_clusters=k_escolhido, n_init=10, random_state=42)
            dados["Cluster"] = kmeans.fit_predict(X_pad)
            print(f"k escolhido: {k_escolhido}")
            display(dados.groupby("Cluster").size().to_frame("n"))
        """),
        cell("markdown", """
            ## 5. Descrever os perfis

            Compare as médias na **escala original**. Nas variáveis ordinais, a média resume os códigos, não anos de idade ou valores monetários. Só agora consultamos `Diabetes_binary` para calcular o percentual observado em cada cluster; ele não mede a qualidade do agrupamento nem a prevalência populacional.
        """),
        cell("code", """
            por_cluster = dados.groupby("Cluster")
            perfis = por_cluster[atributos].mean()
            perfis.insert(0, "n", por_cluster.size())
            perfis["Pré-diabetes/diabetes (%)"] = 100 * por_cluster["Diabetes_binary"].mean()
            display(perfis.round(2))
        """),
        cell("markdown", """
            ## 6. Visualizar com PCA

            O PCA projeta os atributos padronizados em dois componentes. **Os grupos já foram formados no espaço dos atributos; o PCA é usado apenas para visualização.** Observe a variância explicada: a projeção perde parte da informação, e a sobreposição no plano não resume todas as distâncias originais. Variância explicada não é acurácia.

            O segundo gráfico mantém as mesmas coordenadas e cores dos clusters e distingue o indicador observado: círculos = sem diabetes (0); triângulos = pré-diabetes/diabetes (1). A legenda identifica cada combinação de cluster e indicador, com seu número de pessoas. São rótulos da base, não previsões do K-means; o indicador continua fora do agrupamento e do PCA.
        """),
        cell("code", """
            pca = PCA(n_components=2, svd_solver="full")
            coordenadas = pca.fit_transform(X_pad)
            variancia = 100 * pca.explained_variance_ratio_
            print(f"PC1: {variancia[0]:.1f}% | PC2: {variancia[1]:.1f}% | Total: {variancia.sum():.1f}%")

            grupos_pca = sorted(dados["Cluster"].unique())
            cores = {grupo: f"C{indice % 10}" for indice, grupo in enumerate(grupos_pca)}
            fig, eixo = plt.subplots(figsize=(7, 4))
            for grupo in grupos_pca:
                selecao = dados["Cluster"].to_numpy() == grupo
                eixo.scatter(
                    coordenadas[selecao, 0], coordenadas[selecao, 1],
                    s=10, alpha=0.35, color=cores[grupo], label=f"Cluster {grupo}"
                )
            eixo.set(
                title="Projeção dos clusters com PCA",
                xlabel=f"PC1 ({variancia[0]:.1f}%)", ylabel=f"PC2 ({variancia[1]:.1f}%)"
            )
            eixo.legend()
            plt.tight_layout()
            plt.show()

            fig, eixo = plt.subplots(figsize=(11, 5))
            for grupo in grupos_pca:
                for indicador, rotulo, marcador in [
                    (0, "Sem diabetes", "o"),
                    (1, "Pré-diabetes/diabetes", "^"),
                ]:
                    selecao = (
                        (dados["Cluster"].to_numpy() == grupo)
                        & (dados["Diabetes_binary"].to_numpy() == indicador)
                    )
                    eixo.scatter(
                        coordenadas[selecao, 0], coordenadas[selecao, 1],
                        color=cores[grupo], marker=marcador,
                        s=14 if indicador == 0 else 30,
                        alpha=0.25 if indicador == 0 else 0.8,
                        zorder=2 if indicador == 0 else 3,
                        label=f"Cluster {grupo} · {rotulo} (n={selecao.sum():,})"
                    )
            eixo.set(
                title="PCA: clusters e indicador observado de diabetes",
                xlabel=f"PC1 ({variancia[0]:.1f}%)", ylabel=f"PC2 ({variancia[1]:.1f}%)"
            )
            eixo.legend(
                title="Cluster · Indicador da base", loc="upper left",
                bbox_to_anchor=(1.02, 1), fontsize=8, markerscale=1.5
            )
            plt.tight_layout()
            plt.show()
        """),
        cell("markdown", """
            ## Para responder e entregar

            1. Por que padronizar e excluir ID e o indicador de diabetes?
            2. Qual k você escolheu? Justifique usando inércia e silhouette.
            3. Descreva dois clusters com pelo menos três atributos e seus tamanhos. Compare o percentual do indicador de diabetes sem interpretar os grupos como diagnósticos.
            4. Quanta variância os dois componentes preservam? O que limita a leitura do gráfico?
            5. Guarde os resultados iniciais, retire `Income` da lista de atributos e repita as seis células com a mesma semente. Compare k, silhouette, tamanhos e perfis; os números dos clusters podem trocar. As distâncias mudam ao retirar o atributo: um silhouette maior não basta para concluir que a solução é melhor. Se houver ausências, confirme que os IDs da amostra são os mesmos.

            **Entrega:** um `.ipynb` com identificação dos participantes, saídas das duas análises e respostas em Markdown. Duplique as seis células antes de retirar `Income` na cópia, preservando as saídas iniciais. Conclua em 5–8 linhas com um achado e duas limitações. Considere escala, valores extremos, codificação ordinal, amostragem e perda de informação no PCA.

            **Referências:** [K-means](https://scikit-learn.org/stable/modules/generated/sklearn.cluster.KMeans.html), [silhouette](https://scikit-learn.org/stable/modules/generated/sklearn.metrics.silhouette_score.html) e [PCA](https://scikit-learn.org/stable/modules/generated/sklearn.decomposition.PCA.html), na documentação do scikit-learn.
        """),
    ]
    for index, entry in enumerate(cells):
        entry["id"] = f"atividade3-{index:02d}"
    return {
        "cells": cells,
        "metadata": {
            "colab": {"name": NOTEBOOK, "provenance": []},
            "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
            "language_info": {"name": "python"},
        },
        "nbformat": 4,
        "nbformat_minor": 5,
    }


def add_inline(paragraph, text: str) -> None:
    """Converte negrito e links usados no enunciado Markdown para Word."""
    for part in re.split(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))", text):
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            label, url = link.groups()
            if url.startswith("../"):
                url = f"{REPO_URL}/blob/main/{url[3:]}"
            # Espaços e acentos são codificados também no link transportado ao Word.
            add_hyperlink(paragraph, label, quote(unquote(url), safe=":/+"))
        elif part.startswith("**"):
            add_text(paragraph, part[2:-2], bold=True)
        else:
            paragraph.add_run(part)


def build_document(
    task_path: Path = TASK,
    title: str = "Atividade 3 — Aprendizado não supervisionado | Versão mínima",
) -> Document:
    document = Document()
    style_document(document)
    configure_section(document.sections[0])
    document.styles["Title"].font.size = Pt(23)
    document.styles["Heading 1"].font.size = Pt(14)
    document.styles["Heading 2"].font.size = Pt(11)
    for name in ("Normal", "List Bullet"):
        document.styles[name].paragraph_format.keep_together = True
    header = document.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, "APRENDIZADO DE MÁQUINA PARA SAÚDE", bold=True, color=TEAL, size=8)
    add_page_number(document.sections[0].footer.paragraphs[0])
    document.core_properties.title = title
    document.core_properties.language = "pt-BR"

    for block in task_path.read_text(encoding="utf-8").strip().split("\n\n"):
        if block.startswith("#"):
            prefix, title = block.split(" ", 1)
            document.add_heading(title, level=len(prefix) - 1)
        elif block.startswith("- "):
            for line in block.splitlines():
                add_inline(document.add_paragraph(style="List Bullet"), line[2:])
        else:
            add_inline(document.add_paragraph(), block)
    return document


def main() -> None:
    notebook_path = ROOT / "notebooks" / NOTEBOOK
    notebook_path.write_text(
        json.dumps(build_notebook(), ensure_ascii=False, indent=1) + "\n", encoding="utf-8"
    )
    document_path = TASK.with_suffix(".docx")
    build_document().save(document_path)
    print(notebook_path)
    print(document_path)


if __name__ == "__main__":
    main()
